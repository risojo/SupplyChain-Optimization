"""Pruebas de la escritura en el Excel: «fill rate deseado» y pronóstico calculado."""
from __future__ import annotations

import os
import shutil
import sys
import tempfile
import unittest
import zipfile

import numpy as np
import openpyxl
import pandas as pd

_DIR = os.path.dirname(os.path.abspath(__file__))
_PRON = os.path.dirname(_DIR)
if _PRON not in sys.path:
    sys.path.insert(0, _PRON)

import pronostico_calculos as pc  # noqa: E402
import pronostico_excel as pexcel  # noqa: E402


class _ConCarpetaTemporal(unittest.TestCase):
    def setUp(self) -> None:
        self.carpeta = tempfile.mkdtemp()

    def tearDown(self) -> None:
        shutil.rmtree(self.carpeta, ignore_errors=True)


class TestLibroMinimo(_ConCarpetaTemporal):
    def _crear(self) -> str:
        ruta = os.path.join(self.carpeta, "mini.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "data"
        ws.append(["codigo", "clase", "EVAI"])
        for codigo, clase in (("1", "A+"), ("2", "A"), ("3", "B"), ("4", "C"), ("5", "C-"), ("6", "Z")):
            ws.append([codigo, clase, None])
        wb.create_sheet("parametros").append(["x", 1])
        wb.save(ruta)
        return ruta

    def test_agrega_columna_despues_de_la_ultima(self) -> None:
        ruta = self._crear()
        resultado = pexcel.actualizar_fill_rate_excel(ruta)
        self.assertEqual(resultado.estado, "actualizado")
        self.assertEqual(resultado.columna, "D")
        ws = openpyxl.load_workbook(ruta)["data"]
        self.assertEqual(ws["D1"].value, pc.COL_FILL_RATE)
        self.assertEqual([ws[f"D{i}"].value for i in range(2, 8)], [0.997, 0.97, 0.9, 0.6, 0.5, None])
        self.assertEqual(ws["D2"].number_format, "0.0%")

    def test_segunda_pasada_sin_cambios(self) -> None:
        ruta = self._crear()
        pexcel.actualizar_fill_rate_excel(ruta)
        self.assertEqual(pexcel.actualizar_fill_rate_excel(ruta).estado, "sin cambios")

    def test_si_cambia_la_clase_se_actualiza_el_valor(self) -> None:
        ruta = self._crear()
        pexcel.actualizar_fill_rate_excel(ruta)
        wb = openpyxl.load_workbook(ruta)
        wb["data"]["B4"] = "A+"
        wb.save(ruta)
        resultado = pexcel.actualizar_fill_rate_excel(ruta)
        self.assertEqual((resultado.estado, resultado.columna), ("actualizado", "D"))
        self.assertEqual(openpyxl.load_workbook(ruta)["data"]["D4"].value, 0.997)

    def test_archivo_inexistente(self) -> None:
        self.assertEqual(pexcel.actualizar_fill_rate_excel(os.path.join(self.carpeta, "no.xlsx")).estado, "error")


class TestExcelReal(_ConCarpetaTemporal):
    def setUp(self) -> None:
        super().setUp()
        if not os.path.isfile(pc.ARCHIVO_EXCEL_PATH):
            self.skipTest("No está el Excel compartido")
        self.ruta = os.path.join(self.carpeta, "perfilado.xlsx")
        shutil.copy2(pc.ARCHIVO_EXCEL_PATH, self.ruta)

    def test_conserva_formulas_y_resto_de_datos(self) -> None:
        antes = pd.read_excel(self.ruta, pc.HOJA_DATOS).drop(columns=[pc.COL_FILL_RATE], errors="ignore")
        ws_antes = openpyxl.load_workbook(self.ruta)[pc.HOJA_DATOS]
        w2_antes, z2_antes = ws_antes["W2"].value, ws_antes["Z2"].value
        resultado = pexcel.actualizar_fill_rate_excel(self.ruta)
        self.assertIn(resultado.estado, ("actualizado", "sin cambios"))

        despues = pd.read_excel(self.ruta, pc.HOJA_DATOS)
        self.assertEqual(despues.columns[-1], pc.COL_FILL_RATE)
        self.assertEqual(list(despues.columns[-2:]), ["EVAI", pc.COL_FILL_RATE])
        pd.testing.assert_frame_equal(antes, despues.drop(columns=[pc.COL_FILL_RATE]))

        con_clase = despues[despues["clase"].notna()]
        esperado = pc.fill_rate_deseado(con_clase["clase"])
        pd.testing.assert_series_equal(con_clase[pc.COL_FILL_RATE], esperado, check_names=False)

        ws = openpyxl.load_workbook(self.ruta)[pc.HOJA_DATOS]
        self.assertEqual(ws["W2"].value, w2_antes)
        self.assertEqual(ws["Z2"].value, z2_antes)
        self.assertEqual(pd.read_excel(self.ruta, "parametros", header=None).shape,
                         pd.read_excel(pc.ARCHIVO_EXCEL_PATH, "parametros", header=None).shape)

    def test_escribe_pronostico_lri_en_columna_pronostico(self) -> None:
        df, err = pc.cargar_demanda(self.ruta)
        self.assertIsNone(err)
        calc = pc.calcular_pronostico(df)
        antes = pd.read_excel(self.ruta, pc.HOJA_DATOS)
        resultado = pexcel.escribir_pronostico_excel(dict(zip(calc["codigo"], calc[pc.COL_PRONOSTICO_LRI])), self.ruta)
        self.assertIn(resultado.estado, ("actualizado", "sin cambios"))
        self.assertEqual(resultado.columna, "W")

        despues = pd.read_excel(self.ruta, pc.HOJA_DATOS)
        intactas = [c for c in antes.columns if c not in ("pronostico", "pronostico ajustado")]
        pd.testing.assert_frame_equal(antes[intactas], despues[intactas])

        df2, _ = pc.cargar_demanda(self.ruta)
        certificado = calc[pc.COL_CERTIFICACION] == pc.CERTIFICADO
        np.testing.assert_allclose(df2.loc[certificado, pc.COL_PRONOSTICO], calc.loc[certificado, pc.COL_PRONOSTICO_LRI])
        np.testing.assert_allclose(df2.loc[~certificado, pc.COL_PRONOSTICO], calc.loc[~certificado, "pronostico calc"])
        np.testing.assert_allclose(df2.loc[certificado, pc.COL_PRONOSTICO_AJUSTADO], df2.loc[certificado, pc.COL_PRONOSTICO])
        no_cert = df2[~certificado]
        esperado = [pc.pronostico_ajustado(p, s, n) for p, s, n in zip(
            no_cert[pc.COL_PRONOSTICO], no_cert[pc.COL_DESVIACION], no_cert[pc.COL_NIVEL_SERVICIO])]
        np.testing.assert_allclose(no_cert[pc.COL_PRONOSTICO_AJUSTADO], esperado)
        self.assertTrue((pc.calcular_pronostico(df2)[pc.COL_CUADRA_EXCEL] == "Sí").all())

        ws = openpyxl.load_workbook(self.ruta)[pc.HOJA_DATOS]
        codigos = [str(ws[f"A{i}"].value) for i in range(2, ws.max_row + 1)]
        fila_no_cert = int(np.flatnonzero(~certificado.to_numpy())[0])
        n = codigos.index(str(calc.loc[fila_no_cert, "codigo"])) + 2
        self.assertEqual(ws[f"W{n}"].value, f"=AVERAGE(K{n}:V{n})")
        self.assertEqual(ws[f"Z{n}"].value, f"=NORMINV(Y{n},W{n},X{n})")
        fila_cert = int(np.flatnonzero(certificado.to_numpy())[0])
        m = codigos.index(str(calc.loc[fila_cert, "codigo"])) + 2
        self.assertEqual(ws[f"Z{m}"].value, f"=W{m}")
        with zipfile.ZipFile(self.ruta) as z:
            self.assertNotIn("xl/calcChain.xml", z.namelist())
        self.assertEqual(pexcel.escribir_pronostico_excel(dict(zip(calc["codigo"], calc[pc.COL_PRONOSTICO_LRI])), self.ruta).estado,
                         "sin cambios")


_HOJA_COMPARTIDA = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>\
<row r="1"><c r="A1" t="inlineStr"><is><t>codigo</t></is></c><c r="B1" t="inlineStr"><is><t>demanda mes 1</t></is></c>\
<c r="C1" t="inlineStr"><is><t>demanda mes 2</t></is></c><c r="D1" t="inlineStr"><is><t>pronostico</t></is></c>\
<c r="E1" t="inlineStr"><is><t>desviacion standart</t></is></c><c r="F1" t="inlineStr"><is><t>stock seguridad k</t></is></c>\
<c r="G1" t="inlineStr"><is><t>pronostico ajustado</t></is></c></row>\
<row r="2"><c r="A2" t="inlineStr"><is><t>P1</t></is></c><c r="B2"><v>10</v></c><c r="C2"><v>20</v></c>\
<c r="D2"><f t="shared" ref="D2:D4" si="0">AVERAGE(B2:C2)</f><v>15</v></c><c r="E2"><v>5</v></c><c r="F2"><v>0.8</v></c>\
<c r="G2"><f t="shared" ref="G2:G4" si="1">NORMINV(F2,D2,E2)</f><v>19.2</v></c></row>\
<row r="3"><c r="A3"><v>2</v></c><c r="B3"><v>30</v></c><c r="C3"><v>40</v></c>\
<c r="D3"><f t="shared" si="0"/><v>35</v></c><c r="E3"><v>5</v></c><c r="F3"><v>0.8</v></c>\
<c r="G3"><f t="shared" si="1"/><v>39.2</v></c></row>\
<row r="4"><c r="A4" t="inlineStr"><is><t>P3</t></is></c><c r="B4"><v>50</v></c><c r="C4"><v>60</v></c>\
<c r="D4"><f t="shared" si="0"/><v>55</v></c><c r="E4"><v>5</v></c><c r="F4"><v>0.8</v></c>\
<c r="G4"><f t="shared" si="1"/><v>59.2</v></c></row>\
</sheetData></worksheet>"""


class TestEscrituraPronostico(_ConCarpetaTemporal):
    def _crear(self) -> str:
        base = os.path.join(self.carpeta, "base.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "data"
        wb.save(base)
        ruta = os.path.join(self.carpeta, "compartida.xlsx")
        with zipfile.ZipFile(base) as origen, zipfile.ZipFile(ruta, "w") as destino:
            for info in origen.infolist():
                contenido = _HOJA_COMPARTIDA.encode() if info.filename == "xl/worksheets/sheet1.xml" else origen.read(info)
                destino.writestr(info, contenido)
        return ruta

    def test_desplazar_filas(self) -> None:
        self.assertEqual(pexcel._desplazar_filas("AVERAGE(K3:V3)", 4), "AVERAGE(K7:V7)")
        self.assertEqual(pexcel._desplazar_filas("SUM($K$3:K3)+LOG10(B2)", 1), "SUM($K$3:K4)+LOG10(B3)")

    def test_normalizar_codigo(self) -> None:
        self.assertEqual(pexcel.normalizar_codigo(77001046.0), "77001046")
        self.assertEqual(pexcel.normalizar_codigo(" 34-067284 "), "34-067284")

    def test_reemplaza_maestra_compartida_y_conserva_las_demas(self) -> None:
        ruta = self._crear()
        resultado = pexcel.escribir_pronostico_excel({"P1": 50.0, "2": 60.0, "P3": float("nan")}, ruta)
        self.assertEqual((resultado.estado, resultado.columna, resultado.celdas), ("actualizado", "D", 2))

        ws = openpyxl.load_workbook(ruta)["data"]
        self.assertEqual((ws["D2"].value, ws["D3"].value), (50.0, 60.0))
        self.assertEqual(ws["D4"].value, "=AVERAGE(B4:C4)")
        self.assertEqual((ws["G2"].value, ws["G3"].value), ("=D2", "=D3"))
        self.assertEqual(ws["G4"].value, "=NORMINV(F4,D4,E4)")

        datos = pd.read_excel(ruta, "data")
        self.assertEqual(datos.loc[2, "pronostico"], 55)
        self.assertEqual(list(datos["pronostico ajustado"]), [50.0, 60.0, 59.2])
        self.assertEqual(pexcel.escribir_pronostico_excel({"P1": 50.0, "2": 60.0}, ruta).estado, "sin cambios")


if __name__ == "__main__":
    unittest.main()
